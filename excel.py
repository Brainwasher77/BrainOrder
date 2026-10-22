"""Excel de canciones faltantes: nombre, artista, link de Spotify y link de Deezer."""
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config import AUTHOR, COPYRIGHT, CREDIT


def _sign(wb, ws):
    """Marca de agua del autor: propiedades del archivo, encabezado y pie al imprimir."""
    wb.properties.creator = AUTHOR
    wb.properties.lastModifiedBy = AUTHOR
    wb.properties.company = AUTHOR
    wb.properties.description = COPYRIGHT
    ws.oddHeader.right.text = CREDIT
    ws.oddHeader.right.color = "999999"
    ws.oddFooter.center.text = f"{CREDIT} · {COPYRIGHT}"
    ws.oddFooter.center.color = "999999"

HEAD_FILL = PatternFill("solid", fgColor="1C1F27")
HEAD_FONT = Font(bold=True, color="C8F53C", size=11)
TITLE_FONT = Font(bold=True, size=15)
SUB_FONT = Font(italic=True, color="6B7280", size=10)
LINK_SPOTIFY = Font(color="1A7F37", underline="single")
LINK_DEEZER = Font(color="7C3AED", underline="single")
LINK_SEARCH = Font(color="B45309", underline="single", italic=True)
ZEBRA = PatternFill("solid", fgColor="F4F5F7")
LINE = Border(bottom=Side(style="thin", color="E5E7EB"))


def missing_rows(analysis):
    """Canciones que faltan, una sola vez cada una, en el orden en que aparecen."""
    rows = {}
    for pl in analysis:
        for t in pl["tracks"]:
            if t["status"] != "missing":
                continue
            r = rows.get(t["id"])
            if r is None:
                dz = t.get("deezer") or {}
                r = rows[t["id"]] = {
                    "name": t["name"],
                    "artist": ", ".join(t["artists"]),
                    "spotify": t.get("url", ""),
                    "deezer": dz.get("url", ""),
                    "deezer_found": dz.get("found", False),
                    "home": t.get("home_name", pl["name"]),
                    "playlists": [],
                }
            r["playlists"].append(pl["name"])
    return list(rows.values())


def build(analysis, path):
    rows = missing_rows(analysis)
    wb = Workbook()
    ws = wb.active
    ws.title = "Faltantes"

    ws["A1"] = "BrainOrder — Canciones por descargar"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"{len(rows)} canciones · generado {datetime.now():%Y-%m-%d %H:%M} · {CREDIT}"
    ws["A2"].font = SUB_FONT

    headers = ["#", "Canción", "Artista", "Spotify", "Deezer", "Carpeta destino", "Playlists"]
    widths = [6, 42, 30, 48, 42, 26, 40]
    hr = 4
    for c, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=hr, column=c, value=h)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[hr].height = 22

    for i, r in enumerate(rows, 1):
        y = hr + i
        vals = [i, r["name"], r["artist"], r["spotify"], r["deezer"], r["home"], " | ".join(r["playlists"])]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=y, column=c, value=v)
            cell.border = LINE
            cell.alignment = Alignment(vertical="center")
            if i % 2 == 0:
                cell.fill = ZEBRA
        if r["spotify"]:
            ws.cell(row=y, column=4).hyperlink = r["spotify"]
            ws.cell(row=y, column=4).font = LINK_SPOTIFY
        if r["deezer"]:
            cell = ws.cell(row=y, column=5)
            cell.hyperlink = r["deezer"]
            if r["deezer_found"]:
                cell.font = LINK_DEEZER
            else:
                cell.value = f"Buscar en Deezer: {r['artist']} - {r['name']}"
                cell.font = LINK_SEARCH

    ws.freeze_panes = ws.cell(row=hr + 1, column=1)
    if rows:
        ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(headers))}{hr + len(rows)}"
    note = hr + len(rows) + 2
    ws.cell(row=note, column=2, value="Links naranja = Deezer no encontró la misma canción; abren una búsqueda.").font = SUB_FONT
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    _sign(wb, ws)
    wb.save(path)
    return len(rows)


# ---------------- Informe por playlist ----------------
STATUS = {
    "here": ("Aquí", "D1FADF", "166534"),
    "elsewhere": ("En otra carpeta", "DBEAFE", "1E40AF"),
    "inbox": ("En descargas", "FEF3C7", "92400E"),
    "missing": ("Falta", "FEE2E2", "991B1B"),
}


def _state(pl, r):
    from pathlib import Path
    if r["status"] in ("ok", "elsewhere") and r.get("source"):
        return "here" if str(Path(r["source"]).parent) == pl["folder"] else "elsewhere"
    return r["status"] if r["status"] in STATUS else "missing"


def _short(p):
    from pathlib import Path
    home = str(Path.home())
    return p.replace(home, "~", 1) if p else ""


def build_playlist(pl, path):
    """Informe de UNA playlist: todas sus canciones en orden, con estado, ubicación y links."""
    rows = [(r, _state(pl, r)) for r in pl["tracks"]]
    cnt = {k: sum(1 for _, s in rows if s == k) for k in STATUS}
    total = len(rows)
    have = cnt["here"] + cnt["elsewhere"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Informe"
    ws["A1"] = f"{'♥ ' if pl.get('liked') else ''}{pl['name']}"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Informe {CREDIT} · {datetime.now():%Y-%m-%d %H:%M} · carpeta: {_short(pl['folder'])}"
    ws["A2"].font = SUB_FONT

    # Resumen
    summary = [("Canciones", total, None),
               ("Ya las tienes", have, "166534"),
               ("Aquí", cnt["here"], STATUS["here"][2]),
               ("En otra carpeta", cnt["elsewhere"], STATUS["elsewhere"][2]),
               ("En descargas", cnt["inbox"], STATUS["inbox"][2]),
               ("Faltan", cnt["missing"], STATUS["missing"][2]),
               ("Completa", f"{round(have / total * 100) if total else 100}%", "166534")]
    for c, (k, v, color) in enumerate(summary, 2):
        kc = ws.cell(row=4, column=c, value=k)
        kc.font = Font(size=9, color="6B7280", bold=True)
        vc = ws.cell(row=5, column=c, value=v)
        vc.font = Font(size=16, bold=True, color=color or "111111")
        for cell in (kc, vc):
            cell.alignment = Alignment(horizontal="center")

    headers = ["#", "Canción", "Artista", "Estado", "Dónde está / irá", "Spotify", "Deezer", "✓ Descargada"]
    widths = [6, 40, 28, 16, 60, 44, 40, 14]
    hr = 7
    for c, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=hr, column=c, value=h)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[hr].height = 22

    for i, (r, st) in enumerate(rows, 1):
        y = hr + i
        if st in ("here", "elsewhere"):
            where = _short(r["source"])
        elif st == "inbox":
            where = f"{_short(r['source'])}  →  {r.get('home_name', '')}"
        else:
            where = f"irá a: {r.get('home_name', pl['name'])}"
        dz = r.get("deezer") or {}
        vals = [r["pos"] + 1, r["name"], ", ".join(r["artists"]), STATUS[st][0], where,
                r.get("url", ""), dz.get("url", ""), "✓" if st in ("here", "elsewhere") else ""]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=y, column=c, value=v)
            cell.border = LINE
            cell.alignment = Alignment(vertical="center")
        stc = ws.cell(row=y, column=4)
        stc.fill = PatternFill("solid", fgColor=STATUS[st][1])
        stc.font = Font(bold=True, color=STATUS[st][2])
        ws.cell(row=y, column=8).alignment = Alignment(horizontal="center")
        ws.cell(row=y, column=8).font = Font(bold=True, color="166534")
        if r.get("url"):
            ws.cell(row=y, column=6).hyperlink = r["url"]
            ws.cell(row=y, column=6).font = LINK_SPOTIFY
        if dz.get("url"):
            cell = ws.cell(row=y, column=7)
            cell.hyperlink = dz["url"]
            if dz.get("found"):
                cell.font = LINK_DEEZER
            else:
                cell.value = "Buscar en Deezer"
                cell.font = LINK_SEARCH

    ws.freeze_panes = ws.cell(row=hr + 1, column=3)
    if rows:
        ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(headers))}{hr + total}"
    n = hr + total + 2
    ws.cell(row=n, column=2, value="Filtra la columna Estado por \"Falta\" para ver solo lo que te queda por descargar. "
                                   "Links naranja = Deezer no encontró la misma canción; abren una búsqueda.").font = SUB_FONT
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    _sign(wb, ws)
    wb.save(path)
    return {"total": total, "missing": cnt["missing"], "have": have}
